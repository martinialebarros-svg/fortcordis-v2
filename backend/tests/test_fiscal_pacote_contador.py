import io
import json
import os
import sys
import unittest
import zipfile
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from pypdf import PdfReader

BACKEND_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_DIR))
os.environ.setdefault("DATABASE_URL", "sqlite:///:memory:")
os.environ.setdefault("SECRET_KEY", "pacote-contador-test-secret-key-1234567890")

from app.services import fiscal_export_service as service


LOTE = "2a0d3035-8f39-4424-bf66-f922a403f554"
CONFIG = {
    "nome_empresa": "Fort Cordis - configuração sintética",
    "inscricao_estadual": "12345678000195", "inscricao_municipal": "123456",
    "endereco": "Rua do Prestador, 100", "cidade": "Fortaleza", "estado": "CE",
    "telefone": "85999990000", "email": "prestador@example.test",
    "cnae": "7500-1/00", "regime_tributario": "Simples Nacional",
}


def sample_os(os_id=1, **changes):
    row = {
        "os_id": os_id, "numero_os": f"OS-{os_id}",
        "data_atendimento": "2026-08-31T15:00:00-03:00",
        "tipo_cliente": "PJ", "cliente_nome": "Clínica Sintética LTDA",
        "cliente_documento": "12.345.678/0001-95", "clinica_id": 1,
        "cliente_endereco": "Rua do Tomador, 200, sala 1", "cliente_bairro": "Centro",
        "cliente_cidade": "Fortaleza", "cliente_estado": "CE", "cliente_cep": "60000-000",
        "cliente_telefone": "85999991111", "cliente_email": "tomador@example.test",
        "servico_nome": "Ecocardiograma", "status_os": "Pendente",
        "valor_servico": Decimal("250.00"), "valor_desconto": Decimal("0.00"),
        "valor_final": Decimal("250.00"),
        "paciente_nome": "PACIENTE PRIVADO SENTINELA", "tutor_nome": "TUTOR PRIVADO SENTINELA",
        "diagnostico": "DIAGNÓSTICO PRIVADO SENTINELA", "observacoes": "OBSERVAÇÃO PRIVADA SENTINELA",
        "clinica_atividade_cnae": "CNAE_TOMADOR_NAO_COPIAR",
        "descricao_servico": "DESCRIÇÃO CLÍNICA PRIVADA SENTINELA",
    }
    row.update(changes)
    return row


class FiscalPacoteContadorTest(unittest.TestCase):
    def make_package(self, rows=None, **kwargs):
        with patch.object(service, "_get_configuracao", return_value=CONFIG):
            return service.gerar_pacote_contador(
                rows if rows is not None else [sample_os()], None, competencia=kwargs.pop("competencia", "2026-08"),
                lote_id=LOTE, **kwargs,
            )

    def unpack(self, content):
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            return {name: archive.read(name) for name in archive.namelist()}

    def test_zip_contem_planilha_pdf_texto_manifesto_e_dados_de_cadastro(self):
        content, filename = self.make_package()
        self.assertEqual(filename, "FortCordis_2026-08_2a0d3035.zip")
        files = self.unpack(content)
        self.assertEqual(set(files), {"servicos_contador.xlsx", "conferencia_contador.pdf", "mensagem_contador.txt", "manifesto.json"})
        manifesto = json.loads(files["manifesto.json"])
        self.assertEqual(manifesto["versao"], 1)
        self.assertEqual(manifesto["os_ids"], [1])
        self.assertEqual(manifesto["totais"]["valor_final"], "250.00")
        self.assertEqual(manifesto["prestador"]["cnae_cadastrado"], CONFIG["cnae"])
        self.assertEqual(manifesto["tomadores"][0]["cliente_endereco"], "Rua do Tomador, 200, sala 1")
        self.assertTrue(manifesto["gerado_em"].endswith("-03:00"))
        pdf_text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(files["conferencia_contador.pdf"])).pages)
        self.assertIn("Não é nota fiscal", pdf_text)
        self.assertIn("Rua do Tomador", pdf_text)
        self.assertIn("prestador@example.test", pdf_text)
        self.assertIn("31/08/2026", pdf_text)
        self.assertNotIn("NOTA FISCAL DE SERVIÇOS", pdf_text)
        wb = load_workbook(io.BytesIO(files["servicos_contador.xlsx"]), data_only=True)
        self.assertEqual(wb.sheetnames, ["Resumo", "Serviços", "Pendências", "Prestador e lote"])
        self.assertEqual(wb["Serviços"]["D2"].value, "2026-08-31")
        self.assertEqual(wb["Serviços"]["L2"].value, "Pendente")
        self.assertIn("2026-08", files["mensagem_contador.txt"].decode())

    def test_pacote_nao_copia_campos_clinicos_nem_cnae_tomador_ou_iss_presumido(self):
        content, _ = self.make_package(pendencias=[{**sample_os(9), "motivos_pendencia": ["Cadastro incompleto"]}])
        files = self.unpack(content)
        pdf_text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(files["conferencia_contador.pdf"])).pages)
        with zipfile.ZipFile(io.BytesIO(files["servicos_contador.xlsx"])) as archive:
            spreadsheet_xml = "\n".join(archive.read(name).decode(errors="ignore") for name in archive.namelist())
        exported_text = pdf_text + spreadsheet_xml + files["manifesto.json"].decode() + files["mensagem_contador.txt"].decode()
        for private in ["PACIENTE PRIVADO", "TUTOR PRIVADO", "DIAGNÓSTICO PRIVADO", "OBSERVAÇÃO PRIVADA", "DESCRIÇÃO CLÍNICA", "CNAE_TOMADOR_NAO_COPIAR"]:
            self.assertNotIn(private, exported_text)
        self.assertNotIn("aliquota_iss", exported_text)
        self.assertNotIn("valor_iss", exported_text)
        self.assertNotIn("5%", exported_text)

    def test_resumo_agrupa_documento_sem_misturar_pfs_domiciliares(self):
        rows = [sample_os(1, tipo_cliente="PF", cliente_nome="Tomador sintético A", cliente_documento="52998224725", clinica_id=None, clinica_nome="Atendimento domiciliar"),
                sample_os(2, tipo_cliente="PF", cliente_nome="Tomador sintético B", cliente_documento="11144477735", clinica_id=None, clinica_nome="Atendimento domiciliar"),
                sample_os(3, tipo_cliente="PF", cliente_nome="Tomador sintético A", cliente_documento="529.982.247-25", clinica_id=None, clinica_nome="Atendimento domiciliar")]
        files = self.unpack(self.make_package(rows)[0])
        manifesto = json.loads(files["manifesto.json"])
        self.assertEqual(manifesto["quantidade_tomadores"], 2)
        by_doc = {row["cliente_documento"]: row for row in manifesto["tomadores"]}
        self.assertEqual(by_doc["52998224725"]["os_ids"], [1, 3])
        self.assertEqual(by_doc["52998224725"]["totais"]["valor_final"], "500.00")
        self.assertEqual(by_doc["11144477735"]["totais"]["valor_final"], "250.00")

    def test_cnpj_alfanumerico_oficial_preservado_sem_misturar_tomadores(self):
        rows = [sample_os(1, cliente_documento="12.ABC.345/01DE-35"),
                sample_os(2, cliente_documento="00000000E08G12"),
                sample_os(3, cliente_documento=" 12.abc.345/01de-35 ")]
        files = self.unpack(self.make_package(rows, pendencias=[{**sample_os(9, cliente_documento="12.ABC.345/01DE-35"), "motivos_pendencia": ["Conferir cadastro"]}])[0])
        manifesto = json.loads(files["manifesto.json"])
        by_doc = {row["cliente_documento"]: row for row in manifesto["tomadores"]}
        self.assertEqual(set(by_doc), {"12ABC34501DE35", "00000000E08G12"})
        self.assertEqual(by_doc["12ABC34501DE35"]["os_ids"], [1, 3])
        self.assertEqual(by_doc["12ABC34501DE35"]["totais"]["valor_final"], "500.00")
        wb = load_workbook(io.BytesIO(files["servicos_contador.xlsx"]), data_only=True)
        self.assertEqual(wb["Pendências"]["E2"].value, "12ABC34501DE35")
        self.assertIn("12ABC34501DE35", files["mensagem_contador.txt"].decode())
        pdf_text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(files["conferencia_contador.pdf"])).pages)
        self.assertIn("12ABC34501DE35", pdf_text)

    def test_cnpj_alfanumerico_override_equivalente_e_sintaxe_invalida(self):
        files = self.unpack(self.make_package([sample_os(cliente_documento="12.ABC.345/01DE-35")], dados_tomador={"cliente_documento": "12abc34501de35"})[0])
        self.assertEqual(json.loads(files["manifesto.json"])["tomadores"][0]["cliente_documento"], "12ABC34501DE35")
        for document in ["12@ABC34501DE35", "12ABC34501DE3A", "12ÁBC34501DE35"]:
            with self.subTest(document=document), self.assertRaises(ValueError):
                self.make_package([sample_os(cliente_documento=document)])
        with self.assertRaises(ValueError):
            self.make_package([sample_os(tipo_cliente="PF", cliente_documento="5299822472A")])

    def test_documento_invalido_em_pendencia_nao_bloqueia_os_aptas(self):
        files = self.unpack(self.make_package(pendencias=[{**sample_os(9, cliente_documento="12@abc"), "motivos_pendencia": ["Documento inválido"]}])[0])
        wb = load_workbook(io.BytesIO(files["servicos_contador.xlsx"]), data_only=True)
        self.assertEqual(wb["Pendências"]["E2"].value, "12@ABC")
        self.assertEqual(json.loads(files["manifesto.json"])["os_ids"], [1])

    def test_decimais_reconciliam_resumo_detalhes_manifesto_sem_depender_pagamento(self):
        rows = [sample_os(i, valor_servico=Decimal("0.30"), valor_desconto=Decimal("0.10"), valor_final=Decimal("0.20"), status_os="Pago" if i % 2 else "Pendente") for i in range(1, 101)]
        files = self.unpack(self.make_package(rows)[0])
        manifesto = json.loads(files["manifesto.json"])
        self.assertEqual(manifesto["totais"], {"valor_servico": "30.00", "valor_desconto": "10.00", "valor_final": "20.00"})
        wb = load_workbook(io.BytesIO(files["servicos_contador.xlsx"]), data_only=True)
        for field, column, summary_column in [("valor_servico", 9, 13), ("valor_desconto", 10, 14), ("valor_final", 11, 15)]:
            details = sum((Decimal(str(row[column - 1].value)) for row in list(wb["Serviços"])[1:]), Decimal("0"))
            self.assertEqual(details, Decimal(manifesto["totais"][field]))
            self.assertEqual(Decimal(str(wb["Resumo"].cell(2, summary_column).value)), details)

    def test_competencia_local_respeita_limite_utc_e_datas_naive(self):
        rows = [sample_os(1, data_atendimento="2026-09-01T02:59:59Z"), sample_os(2, data_atendimento="2026-08-31T23:59:59")]
        files = self.unpack(self.make_package(rows)[0])
        wb = load_workbook(io.BytesIO(files["servicos_contador.xlsx"]), data_only=True)
        self.assertEqual([wb["Serviços"].cell(row, 4).value for row in (2, 3)], ["2026-08-31", "2026-08-31"])
        with self.assertRaisesRegex(ValueError, "competência"):
            self.make_package([sample_os(data_atendimento="2026-09-01T03:00:00Z")])

    def test_data_geracao_aware_separada_da_competencia_prestacao(self):
        with patch.object(service, "_now_local", return_value=datetime(2026, 10, 14, 18, 0, tzinfo=service.EXPORT_TIMEZONE)):
            manifesto = json.loads(self.unpack(self.make_package()[0])["manifesto.json"])
        self.assertEqual(manifesto["competencia"], "2026-08")
        self.assertEqual(manifesto["gerado_em"], "2026-10-14T18:00:00-03:00")

    def test_pendencias_excluidas_totais_e_campos_nao_whitelisted(self):
        pending = {**sample_os(9, valor_final=Decimal("999.00")), "motivos_pendencia": ["OS cancelada", "Documento ausente"], "elegivel": False}
        files = self.unpack(self.make_package(pendencias=[pending])[0])
        manifesto = json.loads(files["manifesto.json"])
        self.assertEqual(manifesto["quantidade_pendencias"], 1)
        self.assertEqual(manifesto["os_ids"], [1])
        self.assertEqual(manifesto["totais"]["valor_final"], "250.00")
        wb = load_workbook(io.BytesIO(files["servicos_contador.xlsx"]), data_only=True)
        self.assertEqual(wb["Pendências"]["B2"].value, "OS-9")
        self.assertEqual(wb["Pendências"]["F2"].value, "OS cancelada; Documento ausente")

    def test_cabecalho_pendencias_permanece_com_tabela_ao_mudar_pagina(self):
        rows = [sample_os(), sample_os(2), sample_os(3, tipo_cliente="PF", cliente_nome="Tomador sintético A", cliente_documento="52998224725", clinica_id=None)]
        files = self.unpack(self.make_package(rows, pendencias=[{**sample_os(9), "motivos_pendencia": ["Cadastro incompleto"]}])[0])
        pages = [page.extract_text() for page in PdfReader(io.BytesIO(files["conferencia_contador.pdf"])).pages]
        pending_page = next(page for page in pages if "OS-9" in page)
        self.assertIn("Pendências para conferência", pending_page)

    def test_rejeita_duplicatas_canceladas_datas_e_valores_inconsistentes(self):
        scenarios = [[], [sample_os(), sample_os()], [sample_os(status_os="Cancelado")],
                     [sample_os(data_atendimento=None)], [sample_os(os_id=-1)],
                     [sample_os(valor_final=Decimal("100.00"))], [sample_os(valor_servico="NaN")],
                     [sample_os(cliente_documento="")], [sample_os(tipo_cliente="PF")]]
        for rows in scenarios:
            with self.subTest(rows=rows):
                with self.assertRaises(ValueError):
                    self.make_package(rows)
        for competencia in ["2026-13", "2026-8", "202608", "0000-08"]:
            with self.assertRaises(ValueError):
                self.make_package(competencia=competencia)

    def test_override_global_nao_substitui_documentos_distintos(self):
        with self.assertRaisesRegex(ValueError, "documento"):
            self.make_package(dados_tomador={"cliente_documento": "11144477735", "cliente_nome": "Outro tomador"})

    def test_textos_longos_e_markup_sao_legiveis_e_planilha_nao_executa_formula(self):
        row = sample_os(cliente_nome="=Clínica sintética <A&B> " + "Razão social extensa " * 12,
                        cliente_endereco="Rua longa <prédio> & unidade " * 24,
                        servico_nome="Consulta especializada de cardiologia veterinária " * 90)
        files = self.unpack(self.make_package([row])[0])
        wb = load_workbook(io.BytesIO(files["servicos_contador.xlsx"]), data_only=False)
        self.assertEqual(wb["Resumo"]["C2"].data_type, "s")
        self.assertTrue(wb["Resumo"]["C2"].value.startswith("="))
        reader = PdfReader(io.BytesIO(files["conferencia_contador.pdf"]))
        text = "\n".join(page.extract_text() for page in reader.pages)
        self.assertIn("<A&B>", text)
        self.assertIn("Consulta especializada", text)
        self.assertIn("Total do tomador: R$ 250,00", text)
        self.assertGreaterEqual(len(reader.pages), 2)


if __name__ == "__main__":
    unittest.main()
